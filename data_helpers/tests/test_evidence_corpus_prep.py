"""Tests for shared screening and integrated evidence-corpus preparation."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from data_helpers.prep import evidence_corpus_prep as prep
from data_helpers.prep import corpus


SCREENING_COLUMNS = [
    "UT",
    "title",
    "abstract",
    "source",
    "publication_year",
    "wos_categories",
    "doi",
    "s1_r",
    "s2_r",
    "s3_r",
    "s4_r",
    "s1_bio",
    "s2_dir",
    "s3_drivers",
    "s4_link",
]
WOS_UT_A = "WOS:000000000000001"
WOS_UT_B = "WOS:000000000000002"


def _screening_frame() -> pd.DataFrame:
    rows = []
    specs = [
        ("A", True, "ELIGIBLE", 1, 1, 1, 1, "negative"),
        ("B", True, "ELIGIBLE", 1, -1, 1, 1, "mixed"),
        ("C", False, "NOT_ELIGIBLE", 0, 1, 1, 1, "negative"),
        ("D", False, "NOT_ELIGIBLE", 0, 0, 1, 1, None),
    ]
    for ut, eligible, label, s1, s2, s3, s4, direction in specs:
        rows.append(
            {
                "UT": ut,
                "title": f"Title {ut}",
                "authors": "Author",
                "abstract": "Abstract",
                "source": "WOS",
                "publication_year": 2000,
                "wos_categories": "Ecology",
                "doi": f"10/example/{ut}",
                "eligibility": label,
                "s1_r": s1,
                "s2_r": s2,
                "s3_r": s3,
                "s4_r": s4,
                "s1_bio": "biodiversity",
                "s2_dir": direction,
                "s3_drivers": "driver",
                "s4_link": "link",
                "is_eligible": eligible,
                "raw_output": "not retained",
            }
        )
    return pd.DataFrame(rows)


def test_screening_preparation_reconciles_both_grains(tmp_path: Path) -> None:
    source = tmp_path / "screening.csv"
    _screening_frame().to_csv(source, index=False)
    bundle = prep.build_screening_preparation(
        source,
        eligible_columns=SCREENING_COLUMNS,
        chunksize=2,
        repository_root=tmp_path,
    )
    assert bundle.screening["UT"].tolist() == ["A", "B", "C", "D"]
    assert bundle.eligible["UT"].tolist() == ["A", "B"]
    assert "eligibility" in bundle.screening
    assert "eligibility" not in bundle.eligible
    assert "authors" not in bundle.eligible
    assert bundle.manifest["summary"]["eligible_with_unclear_step"] == 1
    assert bundle.manifest["summary"]["eligible_negative"] == 1
    assert bundle.manifest["source"]["path"] == "screening.csv"
    assert bundle.exclusion_overlap["records"].sum() == 2
    assert bundle.exclusion_overlap.iloc[0]["records"] == 1

    store = prep.ScreeningPreparedStore(tmp_path / "prepared")
    paths = store.write(bundle)
    assert {path.name for path in paths} == {
        "screening_publications.parquet",
        "eligible_screening_publications.parquet",
        "screening_exclusion_overlap.csv",
        "manifest.json",
    }
    manifest, overlap = store.load_analysis()
    eligible, _ = store.load_eligible()
    assert manifest["rows"]["screening_publications"] == 4
    assert len(overlap) == 2
    assert eligible["UT"].tolist() == ["A", "B"]
    assert store.remove_eligible_build_cache()
    manifest_after, _ = store.load_analysis()
    assert not store.eligible_path.exists()
    assert manifest_after["consumed_build_cache"]["filename"] == store.eligible_path.name


def test_build_merged_corpus_from_prepared_screening(
    tmp_path: Path, monkeypatch
) -> None:
    eligible = _screening_frame().loc[
        lambda frame: frame["is_eligible"], SCREENING_COLUMNS
    ]
    coding = tmp_path / "driver.csv"
    pd.DataFrame(
        {"UT": ["A", "B"], "driver": ['["Climate Change"]', '["Pollution"]']}
    ).to_csv(coding, index=False)
    monkeypatch.setattr(
        corpus,
        "get_merged_corpus_config",
        lambda: {
            "key_column": "UT",
            "chunksize": 2,
            "screening": {
                "columns": SCREENING_COLUMNS,
                "eligibility_column": "is_eligible",
            },
            "coding_sources": [
                {"name": "driver", "path": coding, "columns": ["driver"]}
            ],
        },
    )
    merged = corpus.build_merged_corpus_from_eligible_screening(eligible)
    assert merged["UT"].tolist() == ["A", "B"]
    assert merged["driver"].tolist() == [
        '["Climate Change"]',
        '["Pollution"]',
    ]


def test_integrated_corpus_normalizes_lists_and_preserves_ut(tmp_path: Path) -> None:
    merged = pd.DataFrame(
        {
            "UT": [WOS_UT_A, WOS_UT_B],
            "title": ["Title A", "Title B"],
            "abstract": ["Abstract A", "Abstract B"],
            "source": ["WOS", "WOS"],
            "publication_year": [2000, 2001],
            "wos_categories": ["Ecology", "Ecology"],
            "doi": ["10/example/A", "10/example/B"],
            "s1_r": [1, 1],
            "s2_r": [1, 1],
            "s3_r": [1, 1],
            "s4_r": [1, 1],
            "s1_bio": ["biodiversity", "biodiversity"],
            "s2_dir": ["negative", "mixed"],
            "s3_drivers": ["driver", "driver"],
            "s4_link": ["link", "link"],
            "pred_study_design": ["Observational", "[]"],
            "driver": ['["Climate Change"]', '["Pollution"]'],
            "pred_threat_l0": [
                '["Climate Change & Severe Weather"]',
                '["Pollution"]',
            ],
            "pred_regions": ['["Americas"]', '[]'],
            "pred_subregions": ['["North America"]', '[]'],
            "pred_countries": ['["USA"]', '[]'],
            "locales": ['["Lake"]', '[]'],
            "locale_coordinates": [None, None],
            "realm": ['["Freshwater"]', '["Terrestrial"]'],
            "pred_methods_data_collection": ['["FieldSurvey"]', '[]'],
            "pred_methods_analysis": ['["DiversityMetrics"]', '[]'],
            "pred_has_comparison": [False, True],
            "pred_comparison_types": ['[]', '["ControlImpact"]'],
        }
    )
    taxa = pd.DataFrame(
        {
            "UT": [WOS_UT_A, WOS_UT_B],
            "publication_year": pd.Series([2000, 2001], dtype="Int64"),
            "s2_dir": ["negative", "mixed"],
            "pred_study_design": ["Observational", "Experimental"],
            "drivers": [("Climate Change",), ("Pollution",)],
            "threat_l0": [
                ("Climate Change & Severe Weather",),
                ("Pollution",),
            ],
            "realms": [("Freshwater",), ("Terrestrial",)],
            "class_labels": [("Aves",), ("Magnoliopsida",)],
            "broad_groups_all": [("Vertebrates",), ("Plants",)],
            "broad_groups": [("Vertebrates",), ("Plants",)],
            "analysis_groups_all": [("Vertebrates",), ("Vascular plants",)],
            "analysis_groups": [("Vertebrates",), ("Vascular plants",)],
            "detail_groups_all": [("Birds",), ("Vascular plants",)],
            "detail_groups": [("Birds",), ("Vascular plants",)],
            "taxa_record_status": ["resolved", "resolved"],
            "n_llm_taxa": [1, 1],
            "n_taxa_matched": [1, 1],
            "n_taxa_unresolved": [0, 0],
            "n_taxa_api_failed": [0, 0],
            "n_drivers": [1, 1],
            "n_broad_groups": [1, 1],
            "taxa_broad_state": ["resolved", "resolved"],
            "n_analysis_groups": [1, 1],
            "taxa_analysis_state": ["resolved", "resolved"],
            "n_detail_groups": [1, 1],
            "taxa_detail_state": ["resolved", "resolved"],
            "taxa_broad_inclusion_state": [
                "Included: at least one benchmarkable broad group",
                "Included: at least one benchmarkable broad group",
            ],
            "match_status_count__exact": [1, 1],
        }
    )
    rank_values = {
        "taxa_kingdom_labels": [("Animalia",), ("Plantae",)],
        "taxa_phylum_labels": [("Chordata",), ("Tracheophyta",)],
        "taxa_class_labels": [("Aves",), ("Magnoliopsida",)],
        "taxa_order_labels": [(), ()],
        "taxa_family_labels": [(), ()],
        "taxa_genus_labels": [(), ()],
        "taxa_species_labels": [(), ()],
    }
    for column, values in rank_values.items():
        taxa[column] = values
    publications, abstracts = prep.build_biodiversity_evidence_corpus(merged, taxa)
    assert publications["id"].tolist() == [1, 2]
    assert publications["UT"].tolist() == [WOS_UT_A, WOS_UT_B]
    assert abstracts[["id", "UT"]].values.tolist() == [
        [1, WOS_UT_A],
        [2, WOS_UT_B],
    ]
    assert publications.loc[0, "pred_countries"] == ("USA",)
    assert publications.loc[0, "pred_countries_audit"] == ("USA",)
    assert publications.loc[0, "pred_study_design"] == ("Observational",)
    assert publications.loc[1, "pred_study_design"] == ("Unclear",)
    assert publications.loc[0, "taxa_class_labels"] == ("Aves",)
    assert publications.loc[0, "taxa_record_status"] == "resolved"
    assert "taxa_summary_json" not in publications
    assert "taxa_matches" not in publications
    assert abstracts.loc[0, "abstract"] == "Abstract A"
    assert "abstract" not in publications
    assert "wos_categories" not in publications
    assert "source" not in publications
    assert "authors" not in publications
    assert "eligibility" not in publications
    assert "taxa_domain_labels" not in publications
    assert "taxa_subkingdom_labels" not in publications
    assert "n_drivers" not in publications
    assert "drivers" not in publications
    assert publications["text_available"].tolist() == [True, True]
    assert publications["plastics_mention"].tolist() == [False, False]

    geography_hierarchy_audit = pd.DataFrame(
        {
            "UT": [WOS_UT_A, WOS_UT_B],
            "status": ["pass_complete", "review"],
            "primary_reason": [
                "complete_paths_valid",
                "no_resolvable_geography",
            ],
            "reasons": [(), ("no_resolvable_geography",)],
            "warnings": [(), ()],
            "pred_regions": [("Americas",), ()],
            "pred_subregions": [("North America",), ()],
            "pred_countries": [("USA",), ()],
        }
    )
    label_summary = prep.apply_geography_hierarchy_audit_labels(
        publications,
        geography_hierarchy_audit,
    )
    assert label_summary.set_index("audit_label_action")[
        "n_publications"
    ].to_dict() == {"preserved": 1, "review_needed": 1}
    assert publications.loc[0, "pred_countries_audit"] == ("USA",)
    assert publications.loc[1, "pred_regions"] == ()
    assert publications.loc[1, "pred_regions_audit"] == (
        "Unclear - Review needed",
    )
    assert publications.loc[1, "pred_subregions_audit"] == (
        "Unclear - Review needed",
    )
    assert publications.loc[1, "pred_countries_audit"] == (
        "Unclear - Review needed",
    )

    source = tmp_path / "source.txt"
    source.write_text("source", encoding="utf-8")
    manifest = prep.build_biodiversity_manifest(
        publications,
        abstracts,
        geography_hierarchy_audit,
        sources={"fake": source},
        repository_root=tmp_path,
    )
    store = prep.BiodiversityEvidenceStore(tmp_path / "corpus")
    paths = store.write(
        prep.BiodiversityEvidenceBuild(
            publications,
            abstracts,
            manifest,
        )
    )
    assert {path.name for path in paths} == {
        "dataset.parquet",
        "dataset.xlsx",
        "dataset_abstracts.parquet",
        "manifest.json",
    }
    loaded = store.load()
    assert loaded.publications.columns.tolist() == list(prep.EVIDENCE_COLUMNS)
    study_design_type = pq.read_schema(store.dataset_path).field(
        "pred_study_design"
    ).type
    assert pa.types.is_list(study_design_type)
    assert study_design_type.value_type == pa.string()
    assert loaded.publications.loc[0, "pred_study_design"] == (
        "Observational",
    )
    assert loaded.publications.loc[0, "pred_countries"] == ("USA",)
    assert loaded.publications.loc[1, "pred_countries_audit"] == (
        "Unclear - Review needed",
    )
    assert loaded.publications.loc[0, "taxa_record_status"] == "resolved"
    assert "taxa_summary_json" not in loaded.publications
    assert "taxa_matches" not in loaded.publications
    assert "rows" not in loaded.manifest
    assert "upstream" not in loaded.manifest
    assert loaded.manifest["sources"]["fake"] == {"path": "source.txt"}
    loaded_abstracts = store.load_abstracts()
    assert loaded_abstracts.loc[0, "abstract"] == "Abstract A"
    assert loaded_abstracts[["id", "UT"]].values.tolist() == [
        [1, WOS_UT_A],
        [2, WOS_UT_B],
    ]
    assert store.validate_xlsx() == {"sheet": "dataset", "columns": 39}
    workbook = load_workbook(store.dataset_xlsx_path, read_only=True, data_only=True)
    worksheet = workbook["dataset"]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    assert len(rows) == len(publications) + 1
    assert list(rows[0]) == list(prep.EVIDENCE_COLUMNS)
    assert rows[1][0] == 1
    assert rows[1][1] == WOS_UT_A
    assert rows[1][list(prep.EVIDENCE_COLUMNS).index("driver")] == '["Climate Change"]'


def test_plastics_text_features_match_case_insensitively_and_require_text() -> None:
    title = pd.Series(
        ["MICROPLASTICS in soil", "Marine debris survey", "Fishing pressure", ""]
    )
    abstract = pd.Series(["", "", "", ""])
    features = prep._compute_plastics_text_features(title, abstract)
    assert features["text_available"].tolist() == [True, True, True, False]
    assert features["plastics_mention"].tolist() == [True, True, False, False]

    # Word boundaries exclude "plasticity"; abstract-only mentions still count.
    title = pd.Series(["Phenotypic plasticity", "Contamination study"])
    abstract = pd.Series(["", "plastic waste in soils"])
    features = prep._compute_plastics_text_features(title, abstract)
    assert features["plastics_mention"].tolist() == [False, True]


def test_publication_ids_are_stable_reversible_and_strict() -> None:
    forward = pd.Series([WOS_UT_A, WOS_UT_B])
    reverse = forward.iloc[::-1]
    forward_ids = prep.derive_publication_ids(forward)
    reverse_ids = prep.derive_publication_ids(reverse)

    assert forward_ids.tolist() == [1, 2]
    assert dict(zip(forward, forward_ids)) == dict(zip(reverse, reverse_ids))
    assert forward_ids.dtype == "int64"
    assert ("WOS:" + forward_ids.map(lambda value: f"{value:015d}")).equals(
        forward
    )

    with pytest.raises(prep.EvidenceCorpusPrepError, match="exactly 15 digits"):
        prep.derive_publication_ids(pd.Series(["A"]))


def test_standardize_geography_lists_preserves_grain_and_columns(
    tmp_path: Path,
) -> None:
    publications = pd.DataFrame(
        {
            "UT": ["A", "B", "C", "D", "E"],
            "pred_regions": [
                (),
                ("Americas", "Unclear"),
                ("Asia and the Pacific",),
                ("Americas",),
                ("All Regions",),
            ],
            "pred_subregions": [
                (),
                ("Unclear", "North America"),
                ("North-East Asia",),
                ("North America",),
                ("All Subregions",),
            ],
            "pred_countries": [
                (),
                ("USA", "Unclear"),
                ("USA", "CHN"),
                ("USA", "TWN"),
                ("Unclear",),
            ],
            "locales": [
                (),
                ("Field site", "Unclear"),
                ("Beijing",),
                ("California",),
                ("Not Applicable",),
            ],
            "locale_coordinates": [
                "[]",
                "[]",
                '[{"locale": "Beijing", "lat": 39.9042, "lon": 116.4074}]',
                None,
                '["Not Applicable"]',
            ],
            "unchanged": [1, 2, 3, 4, 5],
        }
    )
    original_columns = publications.columns.tolist()
    original_uts = publications["UT"].copy()

    audit = prep.standardize_geography_lists(publications)

    assert publications.columns.tolist() == original_columns
    assert publications["UT"].equals(original_uts)
    assert len(publications) == 5
    assert publications["pred_regions"].tolist() == [
        ("Not Applicable",),
        ("Unclear",),
        ("Asia and the Pacific",),
        ("Americas",),
        ("All Regions",),
    ]
    assert publications["pred_subregions"].tolist() == [
        ("Not Applicable",),
        ("Unclear",),
        ("North-East Asia",),
        ("North America",),
        ("Not Applicable",),
    ]
    assert publications["pred_countries"].tolist() == [
        ("Not Applicable",),
        ("Unclear",),
        ("USA", "CHN"),
        ("USA", "TWN"),
        ("Not Applicable",),
    ]
    assert publications["locales"].tolist() == [
        ("Not Applicable",),
        ("Field site", "Unclear"),
        ("Beijing",),
        ("California",),
        ("Not Applicable",),
    ]
    assert publications["locale_coordinates"].tolist() == [
        '["Not Applicable"]',
        '["Not Applicable"]',
        '[{"locale": "Beijing", "lat": 39.9042, "lon": 116.4074}]',
        None,
        '["Not Applicable"]',
    ]
    assert audit.set_index(["column", "rule"])["changed_values"].to_dict() == {
        ("locale_coordinates", "empty_to_not_applicable"): 2,
        ("locales", "empty_to_not_applicable"): 1,
        ("pred_countries", "all_regions_descendant_to_not_applicable"): 1,
        ("pred_countries", "contains_unclear_to_unclear"): 1,
        ("pred_countries", "empty_to_not_applicable"): 1,
        ("pred_regions", "contains_unclear_to_unclear"): 1,
        ("pred_regions", "empty_to_not_applicable"): 1,
        ("pred_subregions", "all_regions_descendant_to_not_applicable"): 1,
        ("pred_subregions", "contains_unclear_to_unclear"): 1,
        ("pred_subregions", "empty_to_not_applicable"): 1,
    }


def test_ipbes_geography_hierarchy_audit_preserves_labels_and_classifies_paths(
    tmp_path: Path,
) -> None:
    mapping_path = tmp_path / "ipbes_regions.json"
    mapping_path.write_text(
        json.dumps(
            {
                "Americas": {
                    "Caribbean": [{"ISO_3166_alpha_3": "ABW"}],
                    "North America": [{"ISO_3166_alpha_3": "USA"}],
                },
                "Asia and the Pacific": {
                    "North-East Asia": [{"ISO_3166_alpha_3": "CHN"}],
                },
                "Europe and Central Asia": {
                    "Central and Western Europe": [
                        {"ISO_3166_alpha_3": "CYP"}
                    ],
                },
            }
        ),
        encoding="utf-8",
    )
    publications = pd.DataFrame(
        {
            "UT": list("ABCDEFGHIJKLM"),
            "pred_regions": [
                ("Americas",),
                ("Americas", "Asia and the Pacific"),
                ("Europe and Central Asia",),
                ("Americas",),
                ("Europe and Central Asia",),
                ("All Regions",),
                ("All Regions",),
                ("Not Applicable",),
                ("Europe and Central Asia",),
                ("Americas",),
                ("Americas", "Asia and the Pacific"),
                ("Americas",),
                ("Americas",),
            ],
            "pred_subregions": [
                ("Caribbean",),
                ("Caribbean", "North-East Asia"),
                ("Central and Western Europe",),
                ("Caribbean",),
                ("Not Applicable",),
                ("Not Applicable",),
                ("North America",),
                ("Not Applicable",),
                ("Not Applicable",),
                ("North America",),
                ("Caribbean",),
                ("Caribbean", "North America"),
                ("North America",),
            ],
            "pred_countries": [
                ("ABW",),
                ("ABW", "CHN"),
                ("Unclear",),
                ("USA",),
                ("Not Applicable",),
                ("Not Applicable",),
                ("Not Applicable",),
                ("Not Applicable",),
                ("CYP",),
                ("?",),
                ("ABW",),
                ("ABW",),
                ("USA", "Not Applicable"),
            ],
        }
    )
    original = publications.copy(deep=True)

    audit = prep.audit_ipbes_geography_hierarchy(
        publications,
        mapping_path=mapping_path,
    )

    assert publications.equals(original)
    assert audit["UT"].tolist() == publications["UT"].tolist()
    assert audit.set_index("UT")["status"].to_dict() == {
        "A": "pass_complete",
        "B": "pass_complete",
        "C": "pass_partial",
        "D": "review",
        "E": "pass_region_only",
        "F": "global",
        "G": "review",
        "H": "no_geography",
        "I": "review",
        "J": "review",
        "K": "pass_complete",
        "L": "pass_complete",
        "M": "review",
    }
    indexed = audit.set_index("UT")
    assert indexed.at["D", "primary_reason"] == "country_path_missing"
    assert indexed.at["G", "primary_reason"] == (
        "all_regions_with_concrete_descendant"
    )
    assert indexed.at["I", "primary_reason"] == (
        "country_present_subregion_non_concrete"
    )
    assert indexed.at["J", "primary_reason"] == "unknown_country_label"
    assert indexed.at["M", "primary_reason"] == (
        "mixed_special_and_concrete_or_global"
    )
    assert indexed.at["K", "warnings"] == ("region_without_listed_subregion",)
    assert indexed.at["L", "warnings"] == (
        "subregion_without_listed_country",
    )
    summary = prep.summarize_geography_hierarchy_audit(audit)
    assert summary["n_publications"].sum() == len(publications)

    for source, target in zip(
        ("pred_regions", "pred_subregions", "pred_countries"),
        prep.GEOGRAPHY_AUDIT_LABEL_COLUMNS,
        strict=True,
    ):
        publications[target] = publications[source].map(tuple)
    source_labels = publications[
        ["pred_regions", "pred_subregions", "pred_countries"]
    ].copy(deep=True)
    label_summary = prep.apply_geography_hierarchy_audit_labels(
        publications,
        audit,
    )
    assert label_summary.set_index("audit_label_action")[
        "n_publications"
    ].to_dict() == {"preserved": 8, "review_needed": 5}
    assert publications[
        ["pred_regions", "pred_subregions", "pred_countries"]
    ].equals(source_labels)
    for ut in ("D", "G", "I", "J", "M"):
        assert publications.loc[
            publications["UT"].eq(ut),
            list(prep.GEOGRAPHY_AUDIT_LABEL_COLUMNS),
        ].iloc[0].tolist() == [
            ("Unclear - Review needed",),
            ("Unclear - Review needed",),
            ("Unclear - Review needed",),
        ]
    assert publications.loc[
        publications["UT"].eq("A"), "pred_countries_audit"
    ].iloc[0] == ("ABW",)
